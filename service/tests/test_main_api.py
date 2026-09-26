import io
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient

from app import main as app_main
from app.main import app

client = TestClient(app)


def _csv_bytes(header: list[str], rows: list[list[str]]) -> bytes:
    lines = [",".join(header)] + [",".join(r) for r in rows]
    return ("\n".join(lines) + "\n").encode("utf-8")


def test_health_still_works():
    resp = client.get("/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


def test_scan_endpoint_runs_deterministic_checks_without_api_keys(monkeypatch):
    """No GROQ_API_KEY/GEMINI_API_KEY set in the test environment — the
    endpoint should still return a full deterministic scorecard, not
    error out, with a warning explaining semantic reasoning was
    skipped."""
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(
        ["email", "amount"],
        [
            ["a@example.com", "100"],
            ["", "200"],  # missing required field
            ["b@example.com", "150"],
        ],
    )
    resp = client.post(
        "/scans",
        files={"file": ("leads.csv", io.BytesIO(csv_content), "text/csv")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["table_name"] == "leads"
    assert body["total_rows"] == 3
    assert body["llm_used"] is False
    assert any("GROQ_API_KEY" in w for w in body["warnings"])
    assert 1 in [r["row_index"] for r in body["flagged_rows"]]  # the blank-email row
    assert body["semantic_iterations"] == 0
    assert body["compilation"] is None


def test_scan_endpoint_with_related_file_enables_referential_integrity(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    accounts_csv = _csv_bytes(["id", "name"], [["1", "Acme"], ["2", "Globex"]])
    leads_csv = _csv_bytes(
        ["email", "account_id"],
        [
            ["a@example.com", "1"],
            ["b@example.com", "9999"],  # references a nonexistent account
        ],
    )

    resp = client.post(
        "/scans",
        files=[
            ("file", ("leads.csv", io.BytesIO(leads_csv), "text/csv")),
            ("related_files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    body = resp.json()
    flagged_indices = {r["row_index"] for r in body["flagged_rows"]}
    assert 1 in flagged_indices
    reasons = next(r for r in body["flagged_rows"] if r["row_index"] == 1)["deterministic_reasons"]
    assert any("referential_integrity_check" in r for r in reasons)


def test_multi_file_scan_request_leaves_no_temp_files_behind(monkeypatch):
    """Regression test for the temp-file leak in `_table_from_upload`:
    every upload (primary + related) is written to a
    `NamedTemporaryFile` and must be unlinked in the `finally` block
    regardless of outcome. This wraps `tempfile.NamedTemporaryFile` as
    called from app.main, records every path it hands back across a
    multi-file /scans request, and asserts none of them survive the
    request — catching a reintroduced leak, not just re-testing that
    the endpoint returns 200."""
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    created_paths: list[str] = []
    real_named_temp_file = tempfile.NamedTemporaryFile

    def _tracking_named_temp_file(*args, **kwargs):
        tmp = real_named_temp_file(*args, **kwargs)
        created_paths.append(tmp.name)
        return tmp

    monkeypatch.setattr(app_main.tempfile, "NamedTemporaryFile", _tracking_named_temp_file)

    accounts_csv = _csv_bytes(["id", "name"], [["1", "Acme"], ["2", "Globex"]])
    leads_csv = _csv_bytes(["email", "account_id"], [["a@example.com", "1"]])

    resp = client.post(
        "/scans",
        files=[
            ("file", ("leads.csv", io.BytesIO(leads_csv), "text/csv")),
            ("related_files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    assert len(created_paths) == 2  # one temp file per uploaded file
    assert all(not Path(p).exists() for p in created_paths)


def test_scan_endpoint_ignores_custom_instruction_without_api_keys(monkeypatch):
    """A custom_instruction can't be compiled or reasoned over without
    an LLM — the endpoint should silently drop it (already covered by
    the warning) rather than crash trying to call a client with no key."""
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(["email", "amount"], [["a@example.com", "100"]])
    resp = client.post(
        "/scans",
        files={"file": ("leads.csv", io.BytesIO(csv_content), "text/csv")},
        data={"custom_instruction": "flag anything weird"},
    )
    assert resp.status_code == 200
    assert resp.json()["compilation"] is None


def test_scan_endpoint_clean_table_scores_100():
    csv_content = _csv_bytes(["email"], [["a@example.com"], ["b@example.com"]])
    resp = client.post(
        "/scans",
        files={"file": ("clean.csv", io.BytesIO(csv_content), "text/csv")},
    )
    body = resp.json()
    assert body["overall_score"] == 100.0
    assert body["flagged_rows"] == []


def test_scan_endpoint_accepts_txt_custom_instruction_file(monkeypatch):
    """No API keys, so the instruction can't actually compile — but the
    endpoint should still read the .txt file's text and pass it through
    as the custom instruction (which then gets silently dropped with the
    existing no-LLM warning, same as a typed instruction would be)."""
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(["email", "amount"], [["a@example.com", "100"]])
    instruction_txt = b"Flag any amount greater than 100000."

    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": ("instructions.txt", io.BytesIO(instruction_txt), "text/plain"),
        },
    )
    assert resp.status_code == 200
    body = resp.json()
    # No LLM configured -> compilation is None either way, but this
    # confirms the file was read and the request didn't error out.
    assert body["compilation"] is None
    assert any("GROQ_API_KEY" in w for w in body["warnings"])


def test_scan_endpoint_accepts_docx_custom_instruction_file(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    import docx

    doc = docx.Document()
    doc.add_paragraph("Flag anything that seems off for this industry.")
    doc.add_paragraph("Pay extra attention to phone numbers.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    csv_content = _csv_bytes(["email", "amount"], [["a@example.com", "100"]])

    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": (
                "instructions.docx",
                buf,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        },
    )
    assert resp.status_code == 200
    assert resp.json()["compilation"] is None


def test_scan_endpoint_rejects_unsupported_instruction_file_type():
    csv_content = _csv_bytes(["email"], [["a@example.com"]])
    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": ("instructions.pdf", io.BytesIO(b"%PDF-1.4"), "application/pdf"),
        },
    )
    assert resp.status_code == 400
    assert "Unsupported" in resp.json()["detail"]


def test_scan_endpoint_rejects_empty_instruction_file():
    csv_content = _csv_bytes(["email"], [["a@example.com"]])
    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": ("instructions.txt", io.BytesIO(b"   \n  "), "text/plain"),
        },
    )
    assert resp.status_code == 400
    assert "readable text" in resp.json()["detail"]


def test_scan_endpoint_truncates_oversized_instruction_file(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(["email"], [["a@example.com"]])
    oversized = ("flag rows that look wrong. " * 300).encode("utf-8")  # well over 4000 chars

    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": ("instructions.txt", io.BytesIO(oversized), "text/plain"),
        },
    )
    assert resp.status_code == 200
    assert any("truncated" in w for w in resp.json()["warnings"])


def test_scan_endpoint_instruction_file_takes_precedence_over_typed_text(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(["email"], [["a@example.com"]])
    resp = client.post(
        "/scans",
        files={
            "file": ("leads.csv", io.BytesIO(csv_content), "text/csv"),
            "custom_instruction_file": ("instructions.txt", io.BytesIO(b"from the file"), "text/plain"),
        },
        data={"custom_instruction": "from the textarea"},
    )
    assert resp.status_code == 200
    warnings = resp.json()["warnings"]
    assert any("typed custom instruction and an instruction file" in w for w in warnings)


def test_schema_scan_requires_at_least_two_tables(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    csv_content = _csv_bytes(["email"], [["a@example.com"]])
    resp = client.post(
        "/scans/schema",
        files=[("files", ("leads.csv", io.BytesIO(csv_content), "text/csv"))],
    )
    assert resp.status_code == 400
    assert "at least 2 tables" in resp.json()["detail"]


def test_schema_scan_runs_per_table_and_flags_cross_table_orphans(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    accounts_csv = _csv_bytes(["id", "name"], [["1", "Acme"], ["2", "Globex"]])
    leads_csv = _csv_bytes(
        ["email", "account_id"],
        [
            ["a@example.com", "1"],
            ["b@example.com", "9999"],  # orphaned FK -> genuine cross-table finding
        ],
    )

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("leads.csv", io.BytesIO(leads_csv), "text/csv")),
            ("files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    body = resp.json()

    assert body["total_tables"] == 2
    assert body["total_rows"] == 4
    assert {t["table_name"] for t in body["tables"]} == {"leads", "accounts"}

    assert len(body["cross_table_findings"]) == 1
    finding = body["cross_table_findings"][0]
    assert finding["from_table"] == "leads"
    assert finding["from_column"] == "account_id"
    assert finding["to_table"] == "accounts"
    assert finding["flagged_row_count"] == 1

    # Aggregate score is a row-count-weighted mean of the two per-table scores.
    leads_score = next(t for t in body["tables"] if t["table_name"] == "leads")["overall_score"]
    accounts_score = next(t for t in body["tables"] if t["table_name"] == "accounts")["overall_score"]
    expected = (leads_score * 2 + accounts_score * 2) / 4
    assert abs(body["aggregate_score"] - round(expected, 2)) < 0.01


def test_schema_scan_with_no_related_tables_has_no_cross_table_findings(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    orders_csv = _csv_bytes(["id", "total"], [["1", "100"], ["2", "200"]])
    employees_csv = _csv_bytes(["id", "name"], [["1", "Alice"], ["2", "Bob"]])

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("orders.csv", io.BytesIO(orders_csv), "text/csv")),
            ("files", ("employees.csv", io.BytesIO(employees_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    assert resp.json()["cross_table_findings"] == []


def test_schema_scan_rejects_malformed_custom_instructions_json():
    accounts_csv = _csv_bytes(["id"], [["1"]])
    leads_csv = _csv_bytes(["email"], [["a@example.com"]])

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("leads.csv", io.BytesIO(leads_csv), "text/csv")),
            ("files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
        data={"custom_instructions": "{not valid json"},
    )
    assert resp.status_code == 400
    assert "JSON object" in resp.json()["detail"]


def test_schema_scan_warns_on_unmatched_instruction_table_name(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    accounts_csv = _csv_bytes(["id"], [["1"]])
    leads_csv = _csv_bytes(["email"], [["a@example.com"]])

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("leads.csv", io.BytesIO(leads_csv), "text/csv")),
            ("files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
        data={"custom_instructions": '{"nonexistent_table": "flag weird rows"}'},
    )
    assert resp.status_code == 200
    assert any("not present in this upload" in w for w in resp.json()["warnings"])


class _FakeGroqForSchemaScan:
    """Stands in for GroqClient — every complete_with_retry call needs
    to answer differently depending on which prompt it's asked to
    handle (compiler / per-table semantic / cross-table semantic), so
    this inspects the prompt content rather than returning one canned
    response for everything."""

    def complete_with_retry(self, system_prompt: str, user_prompt: str) -> str:
        if "linked records across two related tables" in system_prompt:
            # Cross-table semantic pass — flag the contacts/accounts
            # pair whose names visibly don't match.
            if "Totally Different LLC" in user_prompt:
                return (
                    '{"flags": [{"from_row_index": 1, "to_row_index": 1, '
                    '"reason": "contact company name does not match linked account name", '
                    '"confidence": 0.85}]}'
                )
            return '{"flags": []}'
        # No custom instructions used in this test, and autonomous
        # per-table semantic reasoning isn't the point of it either —
        # answer empty for anything else (per-table semantic passes).
        return '{"flags": []}'


class _FakeGeminiForSchemaScan:
    def complete_with_retry(self, system_prompt: str, user_prompt: str) -> str:
        if "pairs of linked records across two tables" in system_prompt:
            return (
                '{"verifications": [{"from_row_index": 1, "to_row_index": 1, '
                '"label": "confirmed", "notes": "names are clearly different companies"}]}'
            )
        return '{"verifications": []}'


def test_schema_scan_runs_cross_table_semantic_pass_with_llm_configured(monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "fake-for-test")
    monkeypatch.setenv("GEMINI_API_KEY", "fake-for-test")
    monkeypatch.setattr("app.agent.llm_clients.GroqClient", lambda: _FakeGroqForSchemaScan())
    monkeypatch.setattr("app.agent.llm_clients.GeminiClient", lambda: _FakeGeminiForSchemaScan())

    accounts_csv = _csv_bytes(["id", "name"], [["1", "Acme Corp"], ["2", "Globex Inc"]])
    contacts_csv = _csv_bytes(
        ["id", "account_id", "company"],
        [
            ["1", "1", "Acme Corp"],  # matches linked account -> no finding
            ["2", "2", "Totally Different LLC"],  # mismatches linked account -> finding
        ],
    )

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("contacts.csv", io.BytesIO(contacts_csv), "text/csv")),
            ("files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["llm_used"] is True

    findings = body["cross_table_semantic_findings"]
    assert len(findings) == 1
    finding = findings[0]
    assert finding["from_table"] == "contacts"
    assert finding["from_row_index"] == 1
    assert finding["to_table"] == "accounts"
    assert finding["to_row_index"] == 1
    assert finding["fk_column"] == "account_id"
    assert finding["verification_label"] == "confirmed"


def test_schema_scan_has_no_cross_table_semantic_findings_without_api_keys(monkeypatch):
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    accounts_csv = _csv_bytes(["id", "name"], [["1", "Acme Corp"]])
    contacts_csv = _csv_bytes(["id", "account_id", "company"], [["1", "1", "Acme Corp"]])

    resp = client.post(
        "/scans/schema",
        files=[
            ("files", ("contacts.csv", io.BytesIO(contacts_csv), "text/csv")),
            ("files", ("accounts.csv", io.BytesIO(accounts_csv), "text/csv")),
        ],
    )
    assert resp.status_code == 200
    assert resp.json()["cross_table_semantic_findings"] == []
